from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt
from reportlab.lib.pagesizes import letter
from reportlab.lib.utils import simpleSplit
from reportlab.pdfgen import canvas

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / 'docs'
PRESENTATION = ROOT / 'presentation'
DOCS.mkdir(exist_ok=True)
PRESENTATION.mkdir(exist_ok=True)


def create_recommendation_pdf(path: Path) -> None:
    c = canvas.Canvas(str(path), pagesize=letter)
    c.setTitle('Executive Recommendation')
    c.setAuthor('Cyncly Case Study')
    c.setFont('Helvetica-Bold', 16)
    c.drawString(72, 740, 'Executive Recommendation')
    c.setFont('Helvetica', 11)
    lines = [
        'Decision: Re-scope and proceed with a controlled migration wave, not a full pause.',
        'Rationale: repository migration is ahead of plan, but pipeline conversion is behind and security controls need stronger enforcement.',
        'Immediate actions: enforce branch protections, separate legacy self-hosted runners, align metric definitions, and reduce developer friction through targeted workflow simplification.',
        'Pause: pause only the migration wave that depends on non-compliant runners or disputed metrics. Do not pause the full transformation.',
        'Security controls that cannot be compromised: least-privilege permissions, OIDC or secret-scoped access, branch protections, and dependency review.',
        'Legacy runners: keep isolated exceptions for unsupported workloads with explicit owner, SLA, and review gates.',
        'Metric reconciliation: align DORA definitions, document exclusions, and require a single authoritative metric dictionary.',
        'Developer friction: measure time-to-first-green, retry rate, and workflow abandonment, then reduce by simplifying PR validation and limiting triggers.',
        'Executive decisions: CTO, Security, Platform Engineering, and Product to approve the re-scoped migration wave and the metric governance model.'
    ]
    y = 700
    for line in lines:
        for wrapped_line in simpleSplit(line, 'Helvetica', 11, 468):
            c.drawString(72, y, wrapped_line)
            y -= 15
        y -= 3
    c.save()


def create_strategy_pdf(path: Path) -> None:
    c = canvas.Canvas(str(path), pagesize=letter)
    c.setTitle('Strategy Package')
    c.setAuthor('Cyncly Case Study')
    c.setFont('Helvetica-Bold', 18)
    c.drawString(72, 760, 'Strategy Package')
    c.setFont('Helvetica', 11)
    lines = [
        'Objective: modernize DevOps tooling and trusted engineering analytics without destabilizing delivery.',
        'Target operating model: GitHub Enterprise, Azure Boards traceability, GitHub Actions validation, and Power BI Desktop reporting.',
        'Platform engineering: reusable workflows, self-hosted runners for legacy workloads, and minimal permission scopes.',
        'Data engineering: local Bronze/Silver/Gold pipeline, explicit quality checks, and Power BI Desktop reporting.',
        '90-day plan: baseline metric definitions, pilot migration, runner review, and governance checkpoints.',
        'Risks: disputed metrics, runner exceptions, and developer friction require executive decision-making and clear trade-offs.',
        'Data caveat: build runs and incident records do not provide verified deployment frequency or event-linked change failure rate.'
    ]
    y = 720
    for line in lines:
        for wrapped_line in simpleSplit(line, 'Helvetica', 11, 468):
            c.drawString(72, y, wrapped_line)
            y -= 16
        y -= 4
    c.save()


def create_presentation(path: Path) -> None:
    prs = Presentation()
    prs.slide_width = 9144000
    prs.slide_height = 5143500

    def add_title(slide, title):
        title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.2), Inches(8.5), Inches(0.5))
        tf = title_box.text_frame
        p = tf.paragraphs[0]
        p.text = title
        p.font.bold = True
        p.font.size = Pt(26)

    def add_bullets(slide, bullets, top=1.1):
        box = slide.shapes.add_textbox(Inches(0.8), Inches(top), Inches(8.5), Inches(4.5))
        tf = box.text_frame
        for i, bullet in enumerate(bullets):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = bullet
            p.level = 0
            p.font.size = Pt(20) if i == 0 else Pt(18)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, 'Executive case and desired outcomes')
    add_bullets(
        slide,
        [
            'Current challenge: repo migration is ahead of plan but pipeline conversion is behind.',
            'Strategic objective: standardize GitHub platform engineering without harming delivery.',
            'Key design principles: least privilege, traceability, quality-first metrics, and controlled coexistence.'
        ]
    )

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, 'Target architecture and operating model')
    add_bullets(
        slide,
        [
            'GitHub Enterprise + Azure Boards traceability for product and work-item governance.',
            'GitHub Actions as validation layer; self-hosted runners reserved for legacy workloads.',
            'Power BI Desktop report over the local Bronze/Silver/Gold DORA data pipeline; deployment measures remain qualified as proxies.'
        ]
    )

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, 'Twelve-month roadmap')
    add_bullets(
        slide,
        [
            'Phase 1: discover and baseline analytics definitions, runner constraints, and developer friction.',
            'Phase 2: pilot GitHub Actions and a controlled DORA metric model on a small service set.',
            'Phase 3: expand to broader migrations with environment approvals and data-quality gates.',
            'Measures: lead time, build-run proxy, MTTR, and limited change-failure proxy with explicit exclusions.'
        ]
    )

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, 'Trusted engineering analytics and AI')
    add_bullets(
        slide,
        [
            'Use a shared metric dictionary to avoid gaming and to reconcile disputed DORA definitions.',
            'Apply quality checks at intake; clearly label metrics as validated, limited, or unavailable.',
            'Use AI for CI diagnosis and test generation with least-privilege access and human approval for production changes.'
        ]
    )

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, 'Executive recommendation and decisions')
    add_bullets(
        slide,
        [
            'Decision: re-scope now rather than pause the full transformation.',
            'Continue: repository migration, metric governance work, and targeted pilot improvements.',
            'Pause: only the migration wave tied to legacy runners or disputed metrics until the controls are approved.',
            'Required owners: CTO, Security, Platform Engineering, Product, and Data/Analytics leadership.'
        ]
    )

    prs.save(path)


create_recommendation_pdf(DOCS / 'EXECUTIVE_RECOMMENDATION.pdf')
create_strategy_pdf(DOCS / 'STRATEGY_PACKAGE.pdf')
create_presentation(PRESENTATION / 'EXECUTIVE_PRESENTATION.pptx')
print('Generated strategy and recommendation PDFs and executive presentation.')
