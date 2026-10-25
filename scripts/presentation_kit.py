"""Reusable design kit for the Cyncly executive presentation.

Generates the brand assets (backgrounds, logo variants, framed screenshot) and
exposes small helpers for shapes and typography so the deck itself reads as
content rather than layout code.
"""

from __future__ import annotations

import random
from pathlib import Path

import numpy as np
from PIL import Image, ImageDraw, ImageFilter, ImageFont
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt
from lxml import etree

ROOT = Path(__file__).resolve().parent.parent
ASSETS = ROOT / "presentation" / "assets"
SCREENSHOT = ROOT / "analytics" / "dashboard" / "dora_powerbi_desktop.png"

# ---------------------------------------------------------------------------
# palette
# ---------------------------------------------------------------------------

INK = "05090F"
BG_DEEP = "060C16"
CARD = "0E1A2C"
CARD_ALT = "111F35"
STROKE = "22375A"
TEXT = "F3F7FD"
TEXT_2 = "A9B9D0"
MUTED = "7386A0"

BLUE = "4A90F4"
BLUE_LT = "7CB6FF"
CYAN = "45CFF5"
VIOLET = "9A5BFF"
INDIGO = "5B5BF0"
MINT = "2FD3A6"
AMBER = "F2B44C"
ROSE = "F0657F"

FONT = "Segoe UI"
FONT_SEMI = "Segoe UI Semibold"
FONT_LIGHT = "Segoe UI Light"

SLIDE_W = 13.333
SLIDE_H = 7.5
MARGIN = 0.62

ALIGN = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT, "j": PP_ALIGN.JUSTIFY}
ANCHOR = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}


def rgb(value: str) -> RGBColor:
    value = value.lstrip("#")
    return RGBColor(int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16))


def _spPr(shape):
    if hasattr(shape, "fill"):
        return shape.fill._xPr
    return shape._element.spPr


def _clear_fill(spPr):
    for tag in ("a:noFill", "a:solidFill", "a:gradFill", "a:blipFill", "a:pattFill", "a:grpFill"):
        el = spPr.find(qn(tag))
        if el is not None:
            spPr.remove(el)


def _insert_after_fill(spPr, element):
    """Insert a fill-adjacent element after fill and before a:ln."""
    ln = spPr.find(qn("a:ln"))
    if ln is not None:
        ln.addprevious(element)
    else:
        spPr.append(element)


def solid(shape, color: str, alpha: int = 100):
    """alpha is opacity in percent."""
    spPr = _spPr(shape)
    _clear_fill(spPr)
    fill = etree.SubElement(spPr, qn("a:solidFill"))
    clr = etree.SubElement(fill, qn("a:srgbClr"))
    clr.set("val", color.lstrip("#"))
    if alpha < 100:
        a = etree.SubElement(clr, qn("a:alpha"))
        a.set("val", str(int(alpha * 1000)))
    _reorder(spPr)


def gradient(shape, stops, angle: float = 90):
    """stops: [(position 0-1, hex, alpha 0-100), ...]."""
    spPr = _spPr(shape)
    _clear_fill(spPr)
    grad = etree.Element(qn("a:gradFill"))
    grad.set("rotWithShape", "1")
    gs_list = etree.SubElement(grad, qn("a:gsLst"))
    for pos, color, alpha in stops:
        gs = etree.SubElement(gs_list, qn("a:gs"))
        gs.set("pos", str(int(pos * 100000)))
        clr = etree.SubElement(gs, qn("a:srgbClr"))
        clr.set("val", color.lstrip("#"))
        if alpha < 100:
            a = etree.SubElement(clr, qn("a:alpha"))
            a.set("val", str(int(alpha * 1000)))
    lin = etree.SubElement(grad, qn("a:lin"))
    lin.set("ang", str(int(angle * 60000)))
    lin.set("scaled", "1")
    _insert_after_fill(spPr, grad)
    _reorder(spPr)


def line_off(shape):
    spPr = _spPr(shape)
    ln = spPr.find(qn("a:ln"))
    if ln is None:
        ln = etree.SubElement(spPr, qn("a:ln"))
    ln.append(etree.Element(qn("a:noFill")))


def outline(shape, color: str, width: float = 0.75, alpha: int = 100, dash: str | None = None):
    spPr = _spPr(shape)
    old = spPr.find(qn("a:ln"))
    if old is not None:
        spPr.remove(old)
    ln = etree.SubElement(spPr, qn("a:ln"))
    ln.set("w", str(int(width * 12700)))
    ln.set("cap", "rnd")
    fill = etree.SubElement(ln, qn("a:solidFill"))
    clr = etree.SubElement(fill, qn("a:srgbClr"))
    clr.set("val", color.lstrip("#"))
    if alpha < 100:
        a = etree.SubElement(clr, qn("a:alpha"))
        a.set("val", str(int(alpha * 1000)))
    if dash:
        d = etree.SubElement(ln, qn("a:prstDash"))
        d.set("val", dash)


def shadow(shape, blur: float = 18, dist: float = 4, direction: int = 5400000,
           color: str = "000000", alpha: int = 40):
    spPr = _spPr(shape)
    effect = spPr.find(qn("a:effectLst"))
    if effect is None:
        effect = etree.Element(qn("a:effectLst"))
        ln = spPr.find(qn("a:ln"))
        if ln is not None:
            ln.addnext(effect)
        else:
            spPr.append(effect)
    shd = etree.SubElement(effect, qn("a:outerShdw"))
    shd.set("blurRad", str(int(blur * 12700)))
    shd.set("dist", str(int(dist * 12700)))
    shd.set("dir", str(direction))
    shd.set("algn", "ctr")
    shd.set("rotWithShape", "0")
    clr = etree.SubElement(shd, qn("a:srgbClr"))
    clr.set("val", color)
    a = etree.SubElement(clr, qn("a:alpha"))
    a.set("val", str(int(alpha * 1000)))


def glow(shape, radius: float = 10, color: str = BLUE, alpha: int = 30):
    spPr = _spPr(shape)
    effect = spPr.find(qn("a:effectLst"))
    if effect is None:
        effect = etree.Element(qn("a:effectLst"))
        ln = spPr.find(qn("a:ln"))
        if ln is not None:
            ln.addnext(effect)
        else:
            spPr.append(effect)
    g = etree.SubElement(effect, qn("a:glow"))
    g.set("rad", str(int(radius * 12700)))
    clr = etree.SubElement(g, qn("a:srgbClr"))
    clr.set("val", color.lstrip("#"))
    a = etree.SubElement(clr, qn("a:alpha"))
    a.set("val", str(int(alpha * 1000)))


def _reorder(spPr):
    """Keep schema order: xfrm, prstGeom/custGeom, fill, ln, effectLst."""
    order = ["a:xfrm", "a:custGeom", "a:prstGeom", "a:noFill", "a:solidFill", "a:gradFill",
             "a:blipFill", "a:pattFill", "a:grpFill", "a:ln", "a:effectLst", "a:effectDag",
             "a:scene3d", "a:sp3d", "a:extLst"]

    def key(el):
        tag = f"a:{etree.QName(el).localname}"
        return order.index(tag) if tag in order else 99

    children = sorted(list(spPr), key=key)
    for child in children:
        spPr.remove(child)
    for child in children:
        spPr.append(child)


# ---------------------------------------------------------------------------
# shapes and typography
# ---------------------------------------------------------------------------

def shape(slide, kind, x, y, w, h, fill=None, *, alpha=100, grad=None, angle=90,
          line=None, lw=0.75, line_alpha=100, dash=None, shade=None, radius=None, rot=0):
    shp = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.shadow.inherit = False
    if radius is not None and kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        shp.adjustments[0] = radius
    if grad:
        gradient(shp, grad, angle)
    elif fill:
        solid(shp, fill, alpha)
    else:
        shp.fill.background()
    if line:
        outline(shp, line, lw, line_alpha, dash)
    else:
        line_off(shp)
    if shade:
        shadow(shp, *shade)
    if rot:
        shp.rotation = rot
    _reorder(_spPr(shp))
    return shp


def card(slide, x, y, w, h, *, radius=0.06, tone=CARD, border=STROKE, alpha=96,
         border_alpha=75, shade=(26, 8, 5400000, "000000", 45)):
    return shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, tone, alpha=alpha,
                 line=border, lw=0.75, line_alpha=border_alpha, radius=radius, shade=shade)


def hairline(slide, x, y, w, color=STROKE, alpha=70, h=0.012):
    shape(slide, MSO_SHAPE.RECTANGLE, x, y, w, h, color, alpha=alpha)


def accent_rule(slide, x, y, w, color=BLUE, h=0.05, alpha=100):
    shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, color, alpha=alpha, radius=0.5)


class Text:
    """Small fluent textbox wrapper."""

    def __init__(self, slide, x, y, w, h, *, anchor="t", margins=(0, 0, 0, 0), wrap=True):
        self.box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = self.box.text_frame
        tf.word_wrap = wrap
        tf.margin_left, tf.margin_right, tf.margin_top, tf.margin_bottom = (Inches(m) for m in margins)
        tf.vertical_anchor = ANCHOR[anchor]
        self.tf = tf
        self.first = True

    def p(self, runs, *, size=12, color=TEXT, bold=False, font=FONT, space_before=0,
          space_after=0, line=1.0, align="l", letter=None, italic=False, opacity=100):
        para = self.tf.paragraphs[0] if self.first else self.tf.add_paragraph()
        self.first = False
        para.alignment = ALIGN[align]
        para.line_spacing = line
        para.space_before = Pt(space_before)
        para.space_after = Pt(space_after)

        if isinstance(runs, str):
            runs = [(runs, {})]
        for chunk, overrides in runs:
            run = para.add_run()
            run.text = chunk
            f = run.font
            f.name = overrides.get("font", font)
            f.size = Pt(overrides.get("size", size))
            f.bold = overrides.get("bold", bold)
            f.italic = overrides.get("italic", italic)
            f.color.rgb = rgb(overrides.get("color", color))
            spc = overrides.get("letter", letter)
            if spc:
                f._rPr.set("spc", str(int(spc * 100)))
            new_opacity = overrides.get("opacity", opacity)
            if new_opacity < 100:
                fill = f._rPr.find(qn("a:solidFill"))
                srgb = fill.find(qn("a:srgbClr"))
                a = etree.SubElement(srgb, qn("a:alpha"))
                a.set("val", str(int(new_opacity * 1000)))
        return para

    def space(self, points=6):
        para = self.tf.paragraphs[0] if self.first else self.tf.add_paragraph()
        self.first = False
        para.space_before = Pt(points)
        para.line_spacing = 1.0
        run = para.add_run()
        run.text = " "
        run.font.size = Pt(1)
        return para


def text(slide, x, y, w, h, content, **kwargs):
    anchor = kwargs.pop("anchor", "t")
    margins = kwargs.pop("margins", (0, 0, 0, 0))
    tb = Text(slide, x, y, w, h, anchor=anchor, margins=margins)
    if isinstance(content, str):
        for line in content.split("\n"):
            tb.p(line, **kwargs)
    else:
        tb.p(content, **kwargs)
    return tb


def pill(slide, x, y, w, h, label, color, *, size=8.5, fill_alpha=14, bold=True,
         text_color=None, border_alpha=45, font=FONT_SEMI):
    shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, color, alpha=fill_alpha,
          line=color, lw=0.75, line_alpha=border_alpha, radius=0.5)
    text(slide, x, y, w, h, label, size=size, color=text_color or color, bold=bold,
         font=font, align="c", anchor="m", letter=0.4)


def badge(slide, x, y, d, label, *, color=BLUE, size=10, fill=None, text_color=None):
    shape(slide, MSO_SHAPE.OVAL, x, y, d, d, fill or color, alpha=100, line=None)
    text(slide, x, y, d, d, label, size=size, color=text_color or INK, bold=True,
         font=FONT_SEMI, align="c", anchor="m")


def logo(slide, x, y, height):
    slide.shapes.add_picture(str(ASSETS / "cyncly-logo-white.png"), Inches(x), Inches(y),
                             height=Inches(height))


# ---------------------------------------------------------------------------
# typography fitting
# ---------------------------------------------------------------------------

_FIT_FONTS = {
    "semibold": Path("C:/Windows/Fonts/seguisb.ttf"),
    "bold": Path("C:/Windows/Fonts/segoeuib.ttf"),
    "regular": Path("C:/Windows/Fonts/segoeui.ttf"),
}
_fit_cache: dict[tuple[str, int], ImageFont.FreeTypeFont] = {}


def text_width(text_value: str, size_pt: float, weight: str = "semibold") -> float:
    """Rendered width in inches, measured from the real font metrics."""
    path = _FIT_FONTS.get(weight)
    if not path or not path.exists():
        return len(text_value) * size_pt * 0.30 / 72
    scale = 4
    key = (str(path), size_pt)
    font = _fit_cache.get(key)
    if font is None:
        font = ImageFont.truetype(str(path), int(size_pt * scale))
        _fit_cache[key] = font
    return font.getlength(text_value) / scale / 72


def fit_size(text_value: str, max_width: float, sizes, weight: str = "semibold") -> float:
    """Largest size from `sizes` that keeps the text on one line."""
    for size in sizes:
        if text_width(text_value, size, weight) <= max_width:
            return size
    return sizes[-1]


# ---------------------------------------------------------------------------
# asset generation
# ---------------------------------------------------------------------------

def _vertical_gradient(w, h, top, bottom):
    top = np.array([int(top[i:i + 2], 16) for i in (0, 2, 4)], dtype=float)
    bottom = np.array([int(bottom[i:i + 2], 16) for i in (0, 2, 4)], dtype=float)
    t = np.linspace(0, 1, h)[:, None, None] ** 0.9
    return top[None, None, :] * (1 - t) + bottom[None, None, :] * t + np.zeros((h, w, 1))


def _add_glow(img, cx, cy, rx, ry, color, strength):
    h, w = img.shape[:2]
    yy, xx = np.mgrid[0:h, 0:w]
    d = np.sqrt(((xx - cx) / rx) ** 2 + ((yy - cy) / ry) ** 2)
    a = np.clip(1 - d, 0, 1) ** 2.4 * strength
    c = np.array([int(color[i:i + 2], 16) for i in (0, 2, 4)], dtype=float)
    return img * (1 - a[..., None]) + c * a[..., None]


def _add_rings(img, cx, cy, radii, color, alpha, width=1.4):
    """Draw soft concentric rings straight into a float RGB canvas."""
    h, w = img.shape[:2]
    yy, xx = np.mgrid[0:h, 0:w]
    dist = np.sqrt((xx - cx) ** 2 + (yy - cy) ** 2)
    c = np.array([int(color[i:i + 2], 16) for i in (0, 2, 4)], dtype=float)
    for r in radii:
        band = np.exp(-((dist - r) ** 2) / (2 * width ** 2))
        img = img * (1 - (band * alpha)[..., None]) + c * (band * alpha)[..., None]
    return img


def build_assets(force=False):
    ASSETS.mkdir(parents=True, exist_ok=True)

    logo_src = ASSETS / "cyncly-logo.png"
    if not logo_src.exists():
        raise FileNotFoundError(
            "presentation/assets/cyncly-logo.png is missing. Download the Cyncly logo first."
        )

    # --- logo variants -----------------------------------------------------
    im = Image.open(logo_src).convert("RGBA")
    w, h = im.size
    split = 105
    white = im.copy()
    px = white.load()
    for yy in range(h):
        for xx in range(split, w):
            r, g, b, a = px[xx, yy]
            px[xx, yy] = (255, 255, 255, a)
    white.save(ASSETS / "cyncly-logo-white.png")

    mark = im.crop((0, 0, 96, 98))
    mark.save(ASSETS / "cyncly-mark.png")
    mono = Image.new("RGBA", mark.size, (255, 255, 255, 0))
    mono.putalpha(mark.getchannel("A").point(lambda a: int(a * 0.85)))
    mono.save(ASSETS / "cyncly-mark-mono.png")

    # --- backgrounds -------------------------------------------------------
    content_path = ASSETS / "bg-content.jpg"
    cover_path = ASSETS / "bg-cover.jpg"
    if force or not content_path.exists():
        W_, H_ = 2560, 1440
        base = _vertical_gradient(W_, H_, "081120", "04080F")
        base = _add_glow(base, W_ * 0.92, -H_ * 0.10, W_ * 0.55, H_ * 0.62, "1B4FA8", 0.42)
        base = _add_glow(base, W_ * 0.06, H_ * 0.98, W_ * 0.42, H_ * 0.55, "431F8F", 0.30)
        base = _add_glow(base, W_ * 0.55, H_ * 0.45, W_ * 0.40, H_ * 0.45, "0D2440", 0.22)
        rgb_img = Image.fromarray(np.clip(base, 0, 255).astype(np.uint8)).convert("RGBA")
        grid = Image.new("RGBA", (W_, H_), (0, 0, 0, 0))
        draw = ImageDraw.Draw(grid)
        for gx in range(0, W_, 120):
            draw.line([(gx, 0), (gx, H_)], fill=(90, 140, 210, 9), width=1)
        for gy in range(0, H_, 120):
            draw.line([(0, gy), (W_, gy)], fill=(90, 140, 210, 9), width=1)
        rgb_img = Image.alpha_composite(rgb_img, grid).convert("RGB")
        rgb_img.save(content_path, quality=92, subsampling=0)

    if force or not cover_path.exists():
        W_, H_ = 2560, 1440
        base = _vertical_gradient(W_, H_, "0A1526", "04070D")
        base = _add_glow(base, W_ * 1.02, H_ * 0.05, W_ * 0.62, H_ * 0.80, "1F5BC8", 0.62)
        base = _add_glow(base, W_ * 0.86, H_ * 0.78, W_ * 0.38, H_ * 0.52, "8B2FE8", 0.34)
        base = _add_glow(base, W_ * 0.40, H_ * 1.05, W_ * 0.55, H_ * 0.45, "0E7FA8", 0.20)
        base = _add_glow(base, W_ * 0.10, H_ * 0.12, W_ * 0.30, H_ * 0.34, "2B3E7A", 0.20)
        base = _add_rings(base, W_ * 1.06, H_ * 0.30, [W_ * 0.30, W_ * 0.40, W_ * 0.51, W_ * 0.63],
                          "6EA8FF", 0.10, width=1.8)
        base = _add_rings(base, W_ * 1.02, H_ * 0.92, [W_ * 0.22, W_ * 0.33, W_ * 0.45],
                          "A970FF", 0.08, width=1.6)

        # faint network of points, right half only
        canvas = Image.fromarray(np.clip(base, 0, 255).astype(np.uint8)).convert("RGBA")
        net = Image.new("RGBA", (W_, H_), (0, 0, 0, 0))
        draw = ImageDraw.Draw(net)
        rnd = random.Random(7)
        pts = [(rnd.randint(int(W_ * 0.58), W_), rnd.randint(0, H_)) for _ in range(34)]
        for i, (x1, y1) in enumerate(pts):
            for x2, y2 in pts[i + 1:]:
                if (x1 - x2) ** 2 + (y1 - y2) ** 2 < (W_ * 0.075) ** 2:
                    draw.line([(x1, y1), (x2, y2)], fill=(125, 170, 245, 20), width=1)
        for x, y in pts:
            draw.ellipse([x - 1.5, y - 1.5, x + 1.5, y + 1.5], fill=(160, 200, 255, 70))
        net = net.filter(ImageFilter.GaussianBlur(0.6))
        canvas = Image.alpha_composite(canvas, net).convert("RGB")
        canvas.save(cover_path, quality=92, subsampling=0)

    # --- dashboard screenshot, cropped and rounded -------------------------
    shot = ASSETS / "dashboard-framed.png"
    if force or not shot.exists():
        src = Image.open(SCREENSHOT).convert("RGB")
        src = src.crop((47, 265, 1227, 822))
        scale = 2
        src = src.resize((src.width * scale, src.height * scale), Image.LANCZOS)
        w2, h2 = src.size
        mask = Image.new("L", (w2, h2), 0)
        ImageDraw.Draw(mask).rounded_rectangle([0, 0, w2 - 1, h2 - 1], radius=26, fill=255)
        out = Image.new("RGBA", (w2, h2), (0, 0, 0, 0))
        out.paste(src, (0, 0), mask)
        ImageDraw.Draw(out).rounded_rectangle([0, 0, w2 - 1, h2 - 1], radius=26,
                                              outline=(40, 70, 110, 255), width=3)
        out.save(shot)


if __name__ == "__main__":
    build_assets(force=True)
    print("assets written to", ASSETS)
