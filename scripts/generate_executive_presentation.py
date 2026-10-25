"""Build the Cyncly CTO executive presentation.

Run:  python scripts/generate_executive_presentation.py
Out:  presentation/EXECUTIVE_PRESENTATION.pptx
"""

from __future__ import annotations

from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from pptx.util import Inches

from presentation_kit import (
    AMBER, ASSETS, BLUE, BLUE_LT, CARD, CARD_ALT, CYAN, FONT, FONT_LIGHT, FONT_SEMI,
    INDIGO, INK, MARGIN, MINT, MUTED, ROSE, SLIDE_H, SLIDE_W, STROKE, TEXT, TEXT_2,
    VIOLET, Text, accent_rule, badge, build_assets, card, fit_size, hairline, logo, pill,
    shadow, shape, text,
)

OUT = Path(__file__).resolve().parent.parent / "presentation" / "EXECUTIVE_PRESENTATION.pptx"
TOTAL = 13


# ---------------------------------------------------------------------------
# deck plumbing
# ---------------------------------------------------------------------------

def add_slide(prs, background="content"):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_picture(str(ASSETS / f"bg-{background}.jpg"), 0, 0,
                             width=prs.slide_width, height=prs.slide_height)
    return slide


def header(slide, idx, part, topic, takeaway, subtitle, accent=BLUE):
    """Two-tone title: the topic in the section colour, the takeaway in white.

    The size is fitted to the available width so a title never wraps into the
    subtitle.
    """
    logo(slide, 11.38, 0.43, 0.32)
    text(slide, MARGIN, 0.44, 7.6, 0.22, part.upper(), size=8.5, color=accent,
         bold=True, font=FONT_SEMI, letter=2.0)

    full = f"{topic} \u2014 {takeaway}"
    size = fit_size(full, 10.45, [25.0, 24.0, 23.0, 21.5, 20.0, 18.5, 17.0])
    tb = Text(slide, MARGIN, 0.68, 11.35, 0.52)
    tb.p([(topic + " \u2014 ", {"color": accent, "size": size}),
          (takeaway, {"color": TEXT, "size": size})],
         size=size, bold=True, font=FONT_SEMI, line=1.0)

    if subtitle:
        text(slide, MARGIN, 1.27, 11.55, 0.44, subtitle, size=10.5, color=TEXT_2, line=1.22)
    hairline(slide, MARGIN, 1.80, SLIDE_W - 2 * MARGIN, alpha=45)


def footer(slide, idx):
    hairline(slide, MARGIN, 7.02, SLIDE_W - 2 * MARGIN, alpha=35)
    text(slide, MARGIN, 7.10, 4.4, 0.2, "Cyncly  ·  Head of DevOps & Analytics", size=7.5,
         color=MUTED)
    text(slide, 5.6, 7.10, 4.2, 0.2, "Internal — for decision", size=7.5, color=MUTED,
         align="c")
    text(slide, 11.3, 7.10, 1.4, 0.2, f"{idx:02d} / {TOTAL:02d}", size=7.5, color=MUTED,
         align="r")


def notes(slide, body):
    slide.notes_slide.notes_text_frame.text = body


def add_transition(slide, kind="fade", speed="med"):
    """Add a slide-show transition; harmless when the deck is edited."""
    tr = etree.SubElement(slide._element, qn("p:transition"))
    tr.set("spd", speed)
    etree.SubElement(tr, qn(f"p:{kind}"))


def arrow(slide, x, y, color=BLUE_LT, size=12):
    text(slide, x, y, 0.3, 0.3, "›", size=size + 8, color=color, bold=True, align="c", anchor="m")


def check_row(slide, x, y, w, label, *, color=MINT, size=9, symbol="✓", gap=0.26):
    badge(slide, x, y + 0.02, gap - 0.04, symbol, color=color, size=7.5, text_color=INK)
    text(slide, x + gap, y, w - gap, 0.3, label, size=size, color=TEXT_2, line=1.05)


def kpi_tile(slide, x, y, w, h, value, unit, label, confidence, conf_color, note):
    card(slide, x, y, w, h, tone=CARD, alpha=98)
    accent_rule(slide, x + 0.02, y + 0.01, w - 0.04, color=conf_color, h=0.045, alpha=85)
    tb = Text(slide, x + 0.2, y + 0.18, w - 0.4, 0.52)
    tb.p([(value, {"size": 27, "font": FONT_LIGHT, "color": TEXT}),
          ("  " + unit, {"size": 10, "color": TEXT_2})], line=1.0)
    text(slide, x + 0.2, y + 0.70, w - 0.4, 0.22, label, size=8.5, color=TEXT_2, bold=True,
         font=FONT_SEMI)
    pill(slide, x + 0.2, y + 1.02, 0.98, 0.24, confidence, conf_color, size=6.8)
    text(slide, x + 1.26, y + 0.99, w - 1.42, 0.38, note, size=7.0, color=MUTED, line=1.08)


# ---------------------------------------------------------------------------
# slide 1 - cover
# ---------------------------------------------------------------------------

def slide_cover(prs):
    s = add_slide(prs, "cover")

    logo(s, 0.88, 0.62, 0.40)
    text(s, 1.23, 1.06, 2.2, 0.2, "Make it amazing", size=7.5, color=TEXT_2, letter=0.6)

    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.88, 2.02, 0.34, 0.028, BLUE, radius=0.5)
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 1.24, 2.02, 0.34, 0.028, VIOLET, radius=0.5)
    text(s, 1.74, 1.92, 8.2, 0.25, "ENGINEERING PLATFORM TRANSFORMATION  ·  25 OCTOBER 2026",
         size=9, color=CYAN, bold=True, font=FONT_SEMI, letter=2.2)

    text(s, 0.88, 2.32, 8.5, 1.9,
         "One platform.\nOne source of truth.\nOne number we can defend.",
         size=38, color=TEXT, bold=True, font=FONT_SEMI, line=1.06)

    text(s, 0.90, 4.42, 7.05, 0.95,
         "Consolidating five Azure DevOps organizations onto GitHub Enterprise — "
         "keeping Azure Boards as the system of record, replacing disputed metrics with "
         "trusted analytics, and applying AI only where it can be governed.",
         size=10.5, color=TEXT_2, line=1.35)

    chip_y = 5.62
    for i, (label, color) in enumerate([("Re-scope, don't pause", MINT),
                                       ("Golden path first", BLUE_LT),
                                       ("Day-90 evidence gate", VIOLET)]):
        pill(s, 0.90 + i * 2.18, chip_y, 2.02, 0.30, label, color, size=8.2, fill_alpha=12)

    hairline(s, 0.88, 6.20, 11.55, alpha=45)

    stats = [("5 → 1", "AZURE DEVOPS ORGANIZATIONS"),
             ("~1,000", "REPOSITORIES"),
             ("~1,500", "PIPELINES"),
             ("12 mo", "TO ONE PLATFORM")]
    for i, (value, label) in enumerate(stats):
        x = 0.88 + i * 2.92
        text(s, x, 6.34, 2.5, 0.38, value, size=21, color=TEXT, font=FONT_LIGHT)
        text(s, x, 6.76, 2.6, 0.2, label, size=7, color=MUTED, bold=True, font=FONT_SEMI, letter=1.4)

    text(s, 0.88, 7.10, 6.0, 0.2, "Head of DevOps & Analytics", size=8, color=TEXT_2, bold=True,
         font=FONT_SEMI)
    text(s, 8.4, 7.10, 4.05, 0.2,
         "Prepared for the CTO and Technology Leadership Team  ·  Internal",
         size=7.5, color=MUTED, align="r")

    notes(s, "Purpose: ask the CTO to re-scope the platform programme, not pause or accelerate it. "
             "The deck has four parts: platform engineering, data engineering, strategy and leadership, "
             "and the executive recommendation. Every figure quoted is reproducible from the supplied "
             "data set and gated by three automated quality checks.")


# ---------------------------------------------------------------------------
# slide 2 - executive summary
# ---------------------------------------------------------------------------

def slide_summary(prs):
    s = add_slide(prs)
    header(s, 2, "The case for change",
           "Executive summary", "migration works; the method does not",
           "Repository migration is ahead of plan — the irreversible work is landing. Pipeline conversion, "
           "control enforcement and measurement are behind, so progress cannot yet be proven.", CYAN)

    cards = [
        ("WHAT IS WORKING", MINT, "Repository migration is ahead of plan",
         ["History and permissions arrive intact",
          "The irreversible work is landing",
          "Credibility earned with delivery teams"]),
        ("WHAT IS STUCK", AMBER, "Pipelines are converted one repository at a time",
         ["~1,500 pipelines, mostly copy-paste YAML",
          "Controls checked quarterly, not enforced",
          "Five organizations, five definitions"]),
        ("WHAT IT COSTS", ROSE, "We cannot prove progress to the board",
         ["No deployment event in the source system",
          "Metrics disputed; MTTR rests on three incidents",
          "Duplication re-created at full scale"]),
    ]
    for i, (eyebrow, color, headline, bullets) in enumerate(cards):
        x = MARGIN + i * 3.98
        card(s, x, 1.98, 3.86, 2.04)
        accent_rule(s, x + 0.02, 1.99, 3.82, color=color, h=0.045, alpha=90)
        text(s, x + 0.22, 2.16, 3.4, 0.2, eyebrow, size=8, color=color, bold=True,
             font=FONT_SEMI, letter=1.6)
        text(s, x + 0.22, 2.42, 3.42, 0.62, headline, size=13.5, color=TEXT, bold=True,
             font=FONT_SEMI, line=1.10)
        tb = Text(s, x + 0.22, 3.10, 3.44, 0.8)
        for b in bullets:
            tb.p([("·  ", {"color": color}), (b, {})], size=8.3, color=TEXT_2,
                 line=1.12, space_after=2.5)

    # recommendation band
    band = shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, MARGIN, 4.16, SLIDE_W - 2 * MARGIN, 1.62,
                 grad=[(0, "101F36", 100), (0.55, "141A36", 100), (1, "1B1436", 100)],
                 angle=0, line="2C4A78", lw=0.75, radius=0.05, shade=(30, 10, 5400000, "000000", 50))
    shape(s, MSO_SHAPE.RECTANGLE, MARGIN, 4.16, 0.07, 1.62,
          grad=[(0, BLUE, 100), (1, VIOLET, 100)], angle=90)

    text(s, MARGIN + 0.30, 4.32, 7.7, 0.2, "THE RECOMMENDATION", size=8, color=VIOLET,
         bold=True, font=FONT_SEMI, letter=2.0)
    text(s, MARGIN + 0.30, 4.54, 8.0, 0.34, "Re-scope — don't accelerate, and don't pause.",
         size=16.5, color=TEXT, bold=True, font=FONT_SEMI)
    text(s, MARGIN + 0.30, 4.96, 7.95, 0.62,
         "Keep migrating. Replace pipeline-by-pipeline conversion with one golden path and a migration "
         "factory. Enforce five controls on new work now. Publish one DORA definition. Release full-estate "
         "funding only on Day-90 evidence.", size=9.2, color=TEXT_2, line=1.28)

    options = [("✕", "NOT ACCELERATE", "Scales the defect, not the fix", ROSE),
               ("✕", "NOT PAUSE", "Stops the only workstream working", ROSE),
               ("✓", "RE-SCOPE", "Moves effort to a reusable platform", MINT)]
    for i, (mark, label, why, color) in enumerate(options):
        y = 4.38 + i * 0.44
        text(s, 9.05, y, 0.24, 0.24, mark, size=9, color=color, bold=True, align="c")
        text(s, 9.30, y - 0.01, 1.35, 0.2, label, size=7.8, color=color, bold=True,
             font=FONT_SEMI, letter=0.8)
        text(s, 10.62, y - 0.01, 2.0, 0.2, why, size=7.6, color=TEXT_2)

    # evidence strip
    card(s, MARGIN, 5.94, SLIDE_W - 2 * MARGIN, 0.82, tone=CARD, alpha=70, border_alpha=70,
         shade=None)
    text(s, MARGIN + 0.24, 6.10, 1.6, 0.5, "EVIDENCE\nALREADY IN HAND", size=7.2, color=MUTED,
         bold=True, font=FONT_SEMI, letter=1.2, line=1.15)
    evidence = [("4 / 4", "unit tests green"), ("3 / 3", "quality gates pass"),
                ("25 / 25", "model files valid"), ("19", "visuals, 2 pages, live")]
    for i, (value, label) in enumerate(evidence):
        x = 2.18 + i * 2.72
        text(s, x, 6.06, 0.85, 0.3, value, size=12.5, color=MINT, bold=True, font=FONT_SEMI)
        text(s, x + 0.80, 6.13, 1.75, 0.3, label, size=7.8, color=TEXT_2)

    footer(s, 2)
    notes(s, "Lead with the asymmetry: repository migration is on plan and irreversible; pipeline conversion "
             "and measurement are the weak links. The recommendation is a re-scope with a 90-day evidence gate. "
             "The evidence strip points to artefacts in the repository the CTO can inspect today.")


# ---------------------------------------------------------------------------
# slide 3 - target architecture
# ---------------------------------------------------------------------------

def slide_architecture(prs):
    s = add_slide(prs)
    header(s, 3, "Part A · Platform engineering",
           "Target architecture", "one org, one path, one metric spine",
           "Five ADO organizations consolidate onto GitHub Enterprise over five waves — distributed "
           "coexistence during migration, one consolidated end state. Azure Boards remains the Agile system "
           "of record in one ADO organization.", BLUE_LT)

    layers = [
        ("1", "Source control", "One GitHub Enterprise org; five ADO orgs retired", "Platform", BLUE),
        ("2", "Change path", "Golden-path workflow inherited, never copy-pasted", "Platform", BLUE),
        ("3", "Execution", "Hosted runners by default; dated exceptions only", "Platform · Security", CYAN),
        ("4", "Telemetry", "Every deployment emits an event; incidents linked at source", "Platform · SRE", CYAN),
        ("5", "Measurement", "Bronze → Silver → Gold → semantic model", "Data engineering", VIOLET),
        ("6", "Consumption", "One baseline report, one definition, one owner", "Data · DevOps", VIOLET),
    ]
    y = 1.98
    for num, name, desc, owner, color in layers:
        card(s, MARGIN, y, 7.58, 0.63, tone=CARD, alpha=92, shade=None)
        badge(s, MARGIN + 0.14, y + 0.17, 0.29, num, color=color, size=8.5, text_color=INK)
        text(s, MARGIN + 0.55, y + 0.11, 5.0, 0.22, name, size=9.8, color=TEXT, bold=True,
             font=FONT_SEMI)
        text(s, MARGIN + 0.55, y + 0.34, 5.6, 0.2, desc, size=7.9, color=TEXT_2)
        pill(s, MARGIN + 6.22, y + 0.19, 1.24, 0.25, owner, color, size=7.0, fill_alpha=10)
        y += 0.71

    # end-state card
    card(s, 8.43, 1.98, 4.28, 2.08)
    text(s, 8.63, 2.12, 3.8, 0.2, "END STATE AT MONTH 12", size=7.5, color=CYAN, bold=True,
         font=FONT_SEMI, letter=1.6)

    text(s, 8.63, 2.38, 1.4, 0.18, "TODAY", size=6.5, color=MUTED, bold=True, font=FONT_SEMI, letter=1.2)
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 8.63, 2.56, 1.42, 0.74, ROSE, alpha=9,
          line=ROSE, lw=0.75, line_alpha=40, radius=0.08)
    text(s, 8.63, 2.56, 1.42, 0.74, "5\nADO organizations", size=8.2, color=TEXT, bold=True,
         font=FONT_SEMI, align="c", anchor="m", line=1.05)
    text(s, 10.10, 2.80, 0.30, 0.25, "›", size=16, color=CYAN, bold=True, align="c")

    text(s, 10.45, 2.38, 1.9, 0.18, "TARGET", size=6.5, color=MUTED, bold=True, font=FONT_SEMI, letter=1.2)
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 10.45, 2.56, 2.04, 0.74, MINT, alpha=9,
          line=MINT, lw=0.75, line_alpha=45, radius=0.08)
    text(s, 10.45, 2.56, 2.04, 0.74, "1\nGitHub Enterprise org", size=8.2, color=TEXT, bold=True,
         font=FONT_SEMI, align="c", anchor="m", line=1.05)

    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 8.63, 3.42, 3.86, 0.50, INDIGO, alpha=10,
          line=INDIGO, lw=0.75, line_alpha=45, radius=0.1)
    text(s, 8.78, 3.42, 3.6, 0.50,
         "Azure Boards — retained in one ADO org\nas the Agile system of record",
         size=7.8, color=TEXT_2, line=1.15, anchor="m")

    # non-negotiables
    card(s, 8.43, 4.18, 4.28, 2.13)
    text(s, 8.63, 4.32, 3.8, 0.2, "NON-NEGOTIABLE BOUNDARIES", size=7.5, color=VIOLET, bold=True,
         font=FONT_SEMI, letter=1.6)
    rules = ["OIDC federation only — no long-lived cloud credentials",
             "Policy enforced as code in the same pull request",
             "Every deployment emits service, environment and commit",
             "The measurement layer reads from one Gold source"]
    yy = 4.58
    for rule in rules:
        shape(s, MSO_SHAPE.OVAL, 8.63, yy + 0.07, 0.075, 0.075, VIOLET)
        text(s, 8.78, yy, 3.72, 0.34, rule, size=8.2, color=TEXT_2, line=1.1)
        yy += 0.34
    text(s, 8.63, 5.98, 3.78, 0.24,
         "A service can request a dated exception — never buy its way around it.",
         size=7.2, color=MUTED, italic=True, line=1.05)

    footer(s, 3)
    notes(s, "Six layers, each with a single owner. The migration target is one GitHub organization, but Azure "
             "Boards stays where work-item governance lives — that decision is deliberate and avoids reopening "
             "traceability design. The four boundaries are the ones we will not negotiate per team.")


# ---------------------------------------------------------------------------
# slide 4 - golden path and controls
# ---------------------------------------------------------------------------

def slide_golden_path(prs):
    s = add_slide(prs)
    header(s, 4, "Part A · Platform engineering",
           "Golden path", "1,500 pipelines converted once",
           "Per-repository porting reproduces the duplication that made the estate expensive. A service inherits "
           "the golden path; it cannot fork its way around it.", BLUE_LT)

    steps = [("Commit", "branch + reviewable change"),
             ("Pull request", "policy and owner checks"),
             ("Automated gates", "tests, SAST, deps, secrets"),
             ("Golden path", "one inherited workflow"),
             ("Deployment event", "service, env, commit"),
             ("Metric grain", "project × month")]
    x = 0.66
    for i, (title, sub) in enumerate(steps):
        card(s, x, 1.98, 1.72, 1.02, tone=CARD_ALT, alpha=95, shade=None,
             border=BLUE if i == 3 else STROKE)
        text(s, x + 0.15, 2.08, 1.4, 0.18, f"0{i + 1}", size=7, color=BLUE_LT, bold=True,
             font=FONT_SEMI, letter=1.2)
        text(s, x + 0.15, 2.28, 1.46, 0.24, title, size=9.2, color=TEXT, bold=True,
             font=FONT_SEMI, line=1.0)
        text(s, x + 0.15, 2.55, 1.46, 0.36, sub, size=7.0, color=TEXT_2, line=1.08)
        if i < 5:
            arrow(s, x + 1.72, 2.33, CYAN, size=10)
        x += 1.94

    # five controls
    text(s, MARGIN, 3.28, 5.9, 0.25, "Five controls — non-waivable on new work", size=11,
         color=TEXT, bold=True, font=FONT_SEMI)
    controls = [
        "Reviewed ruleset on the default branch; required checks green",
        "Third-party actions pinned to full commit SHAs",
        "Workload identity federation; no long-lived secrets",
        "Attributable human approval on production deploys",
        "Secret scanning and push protection on; secrets revoked, not deleted",
    ]
    yy = 3.62
    for c in controls:
        check_row(s, MARGIN, yy, 6.0, c, color=MINT, size=8.6)
        yy += 0.44
    card(s, MARGIN, 5.90, 5.98, 0.78, tone=AMBER, alpha=7, border=AMBER, border_alpha=35,
         shade=None)
    text(s, MARGIN + 0.18, 5.98, 5.62, 0.62,
         "Existing estate: dated backlog, 3-day exception turnaround, maximum 90 days — "
         "renewable once with justification, then closed.",
         size=7.8, color=TEXT_2, line=1.18, anchor="m")

    # runner strategy
    text(s, 7.02, 3.28, 5.7, 0.25, "Runner strategy: standard by default", size=11,
         color=TEXT, bold=True, font=FONT_SEMI)
    runners = [("Standard hosted", "All new and migrated pipelines", "Permanent baseline", MINT),
               ("Larger hosted", "Heavy builds; cost attributed to product", "Reviewed quarterly", BLUE_LT),
               ("Self-hosted", "Only legacy products that cannot move", "≤ 12 months, dated", AMBER)]
    yy = 3.62
    for name, use, control, color in runners:
        card(s, 7.02, yy, 5.69, 0.56, tone=CARD, alpha=90, shade=None)
        shape(s, MSO_SHAPE.RECTANGLE, 7.02, yy, 0.05, 0.56, color, alpha=90)
        text(s, 7.18, yy + 0.08, 2.0, 0.2, name, size=8.8, color=TEXT, bold=True, font=FONT_SEMI)
        text(s, 7.18, yy + 0.30, 2.7, 0.2, use, size=7.2, color=TEXT_2)
        text(s, 9.95, yy + 0.19, 1.55, 0.2, control, size=7.4, color=color, bold=True,
             font=FONT_SEMI, align="r")
        yy += 0.66

    text(s, 7.02, 5.62, 5.7, 0.2, "EXCEPTION PATH", size=7, color=MUTED, bold=True,
         font=FONT_SEMI, letter=1.4)
    path = [("Request", "same day"), ("Risk assess", "2 days"), ("CISO approval", "1 day"), ("Review", "monthly")]
    x = 7.02
    for i, (step, when) in enumerate(path):
        pill(s, x, 5.86, 1.28, 0.30, step, VIOLET, size=7.2, fill_alpha=12)
        text(s, x, 6.20, 1.28, 0.18, when, size=6.5, color=MUTED, align="c")
        if i < 3:
            text(s, x + 1.26, 5.88, 0.18, 0.24, "›", size=10, color=VIOLET, align="c")
        x += 1.44

    footer(s, 4)
    notes(s, "The flow is the contract: each arrow has an owner, and a missing arrow is a finding rather than "
             "something a team works around locally. Five controls are enforced in the pull request, not in a "
             "quarterly audit. Runners are hosted by default; the only exceptions are named, dated and funded.")


# ---------------------------------------------------------------------------
# slide 5 - AI-enabled engineering
# ---------------------------------------------------------------------------

def slide_ai(prs):
    s = add_slide(prs)
    header(s, 5, "Part A · Platform engineering",
           "AI-enabled engineering", "four uses, four enforced guardrails",
           "Copilot in GitHub Enterprise, applied to work that is repetitive, reviewable and measurable. "
           "No customer data in prompts; no autonomous production change.", CYAN)

    uses = [
        ("Pipeline migration copilot", BLUE_LT,
         "Translate ADO YAML into golden-path workflows. The 1,500-pipeline estate is the use case.",
         "Human review and policy check; no direct push"),
        ("Code review and test generation", CYAN,
         "PR summaries, missing-test detection, dependency upgrades, reviewable suggestions.",
         "Suggestions only; human approval stays required"),
        ("Incident triage and postmortems", VIOLET,
         "Correlate telemetry, draft the timeline, propose follow-ups with owners.",
         "Read-only access; no personal data to a model"),
        ("Analytics co-pilot", MINT,
         "Natural-language questions over the semantic model, never over raw tables.",
         "Semantic layer only; every answer reproducible"),
    ]
    for i, (title, color, body, guard) in enumerate(uses):
        x = MARGIN + (i % 2) * 3.55
        y = 1.98 + (i // 2) * 1.72
        card(s, x, y, 3.42, 1.60)
        shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x + 0.18, y + 0.18, 0.34, 0.05, color, radius=0.5)
        text(s, x + 0.18, y + 0.30, 3.05, 0.24, title, size=9.8, color=TEXT, bold=True,
             font=FONT_SEMI, line=1.0)
        text(s, x + 0.18, y + 0.58, 3.06, 0.55, body, size=7.8, color=TEXT_2, line=1.16)
        text(s, x + 0.18, y + 1.16, 3.06, 0.36,
             [("GUARDRAIL  ", {"color": color, "size": 6.6, "bold": True, "letter": 1.0}),
              (guard, {"color": MUTED, "size": 7.1})], line=1.1)

    # measurement panel
    card(s, 7.93, 1.98, 4.78, 3.32)
    text(s, 8.14, 2.14, 4.3, 0.22, "HOW WE PROVE IT PAYS", size=7.5, color=CYAN, bold=True,
         font=FONT_SEMI, letter=1.6)
    measures = [("Adoption", "Share of active repositories using AI-assisted pull requests"),
                ("Time to first green", "Under one hour for a new service on the golden path"),
                ("DORA delta", "Lead time and change failure rate, always reported as a pair"),
                ("Review load", "Comments per pull request — quality, not volume")]
    yy = 2.44
    for name, desc in measures:
        text(s, 8.14, yy, 1.35, 0.2, name, size=8.0, color=TEXT, bold=True, font=FONT_SEMI)
        text(s, 9.48, yy, 3.05, 0.3, desc, size=7.4, color=TEXT_2, line=1.05)
        hairline(s, 8.14, yy + 0.42, 4.38, alpha=18)
        yy += 0.55
    text(s, 8.14, 4.74, 4.32, 0.42,
         "Measured monthly. A negative DORA result stops the rollout — capability is not a strategy.",
         size=7.4, color=MUTED, italic=True, line=1.15)

    # sequence strip
    card(s, MARGIN, 5.50, SLIDE_W - 2 * MARGIN, 1.16, tone=CARD, alpha=70, shade=None,
         border_alpha=70)
    text(s, MARGIN + 0.24, 5.64, 2.2, 0.2, "SEQUENCE, NOT SPRAY", size=7.2, color=MUTED, bold=True,
         font=FONT_SEMI, letter=1.4)
    stages = [("1  Assistant", "autocomplete, chat, docs — lowest risk"),
              ("2  Reviewer", "PR summaries, tests, dependency fixes"),
              ("3  Agent", "scoped migration tasks with a named owner")]
    for i, (label, desc) in enumerate(stages):
        x = MARGIN + 0.24 + i * 3.95
        text(s, x, 5.88, 1.9, 0.2, label, size=9.0, color=CYAN if i == 2 else TEXT, bold=True,
             font=FONT_SEMI)
        text(s, x, 6.12, 3.55, 0.22, desc, size=7.4, color=TEXT_2)
        if i < 2:
            text(s, x + 3.62, 5.90, 0.3, 0.25, "›", size=13, color=BLUE_LT, bold=True, align="c")

    footer(s, 5)
    notes(s, "AI is applied where the work is repetitive and the output is reviewable. Four use cases, each with "
             "a guardrail that is enforced by the platform rather than by a policy document. Sequence matters: "
             "assistant, then reviewer, then bounded agent tasks. Adoption is measured alongside DORA, and a "
             "negative result stops the rollout.")


# ---------------------------------------------------------------------------
# slide 6 - measurement spine
# ---------------------------------------------------------------------------

def slide_data_spine(prs):
    s = add_slide(prs)
    header(s, 6, "Part B · Data engineering",
           "Measurement spine", "a gated pipeline, not a dashboard",
           "One definition, computed once in Gold and published once. Every figure recomputes from stored facts — "
           "and states how far it can be trusted.", VIOLET)

    stages = [
        ("Bronze", "Raw workbook preserved", "SHA-256 + row counts", BLUE),
        ("Silver", "Typed, validated, deduplicated", "Rejects kept, not dropped", CYAN),
        ("Gold", "Project × calendar month", "Four metrics, 9 grain rows", MINT),
        ("Semantic model", "8 tables · 8 relationships", "17 measures, 1 definition", VIOLET),
        ("Report", "Power BI, 2 pages", "19 visuals, refresh-verified", AMBER),
    ]
    x = MARGIN
    for i, (name, desc, chip, color) in enumerate(stages):
        card(s, x, 1.98, 2.25, 1.44, tone=CARD_ALT, alpha=95, shade=None)
        text(s, x + 0.18, 2.10, 1.9, 0.18, f"STAGE {i + 1}", size=6.4, color=color, bold=True,
             font=FONT_SEMI, letter=1.3)
        text(s, x + 0.18, 2.32, 1.95, 0.24, name, size=10.5, color=TEXT, bold=True, font=FONT_SEMI)
        text(s, x + 0.18, 2.60, 1.95, 0.34, desc, size=7.4, color=TEXT_2, line=1.1)
        text(s, x + 0.18, 3.04, 1.95, 0.24, chip, size=6.8, color=color, bold=True, font=FONT_SEMI)
        if i < 4:
            arrow(s, x + 2.25, 2.58, VIOLET, size=9)
        x += 2.44

    text(s, MARGIN, 3.60, 7.7, 0.22, "Three gates, executable, and proven able to fail", size=10.5,
         color=TEXT, bold=True, font=FONT_SEMI)
    gates = [("DQ-01", "Referential integrity and duplicate grain",
              "No orphan facts; one row per project-month", BLUE_LT),
             ("DQ-02", "Temporal validity",
              "No negative lead time, no build before commit", CYAN),
             ("DQ-03", "Metric reconciliation",
              "Published metrics recompute from the facts", MINT)]
    for i, (gid, title, desc, color) in enumerate(gates):
        x = MARGIN + i * 2.50
        card(s, x, 3.90, 2.38, 1.28, tone=CARD, alpha=92, shade=None)
        text(s, x + 0.16, 4.02, 1.2, 0.2, gid, size=7.2, color=color, bold=True, font=FONT_SEMI,
             letter=1.2)
        pill(s, x + 1.68, 3.99, 0.54, 0.21, "PASS", MINT, size=6.5, fill_alpha=16)
        text(s, x + 0.16, 4.32, 2.06, 0.36, title, size=8.4, color=TEXT, bold=True, font=FONT_SEMI,
             line=1.05)
        text(s, x + 0.16, 4.74, 2.06, 0.36, desc, size=7.0, color=TEXT_2, line=1.1)

    card(s, MARGIN, 5.32, 7.42, 1.40, tone=CARD, alpha=60, shade=None, border_alpha=60)
    text(s, MARGIN + 0.22, 5.44, 7.0, 0.22, "PROOF, NOT ASSERTION", size=7.2, color=MINT, bold=True,
         font=FONT_SEMI, letter=1.4)
    proofs = ["Negative test corrupts Gold and requires all three gates to fail — they do.",
              "25 / 25 definition files validate against Microsoft schemas; every visual field resolves.",
              "345 builds, 149 lead-time links and 3 incidents reconcile row by row."]
    yy = 5.68
    for p in proofs:
        text(s, MARGIN + 0.22, yy, 7.0, 0.2, "·  " + p, size=7.5, color=TEXT_2, line=1.05)
        yy += 0.28

    card(s, 8.30, 3.60, 4.41, 3.12)
    text(s, 8.50, 3.74, 4.0, 0.2, "RULES THAT KEEP NUMBERS HONEST", size=7.4, color=VIOLET,
         bold=True, font=FONT_SEMI, letter=1.4)
    rules = [("One owner, one definition", "A change to formula, grain or exclusion is a version bump."),
             ("Pairs, not singles", "Frequency always travels with failure rate."),
             ("Sample sizes published", "A median over three points is labelled indicative."),
             ("Blank is not zero", "No observation stays blank; zero-filling is a fiction.")]
    yy = 4.04
    for title, desc in rules:
        text(s, 8.50, yy, 4.0, 0.2, title, size=8.4, color=TEXT, bold=True, font=FONT_SEMI)
        text(s, 8.50, yy + 0.20, 4.0, 0.3, desc, size=7.2, color=TEXT_2, line=1.1)
        yy += 0.60

    footer(s, 6)
    notes(s, "The data product is the pipeline: raw preserved in Bronze, validated in Silver with rejects kept, "
             "published at one grain in Gold, and modelled once for consumption. The three gates run in CI and a "
             "negative test proves they can fail. The right-hand rules are what prevent a metric from being "
             "improved by mislabelling rather than by doing the work.")


# ---------------------------------------------------------------------------
# slide 7 - trusted metrics and limits
# ---------------------------------------------------------------------------

def slide_metrics(prs):
    s = add_slide(prs)
    header(s, 7, "Part B · Data engineering",
           "Evidence", "four DORA metrics and their honest limits",
           "Recomputed from source, reconciled row by row, gated by three automated checks. Where a number is a "
           "proxy, it is labelled a proxy.", VIOLET)

    kpi_tile(s, MARGIN, 1.98, 2.55, 1.50, "38.33", "", "Deployments per project-month",
             "PROXY", AMBER, "345 builds · 9 grain rows")
    kpi_tile(s, 3.34, 1.98, 2.55, 1.50, "20.5", "h", "Lead time for changes (median)",
             "HIGH", MINT, "n = 149 · 98.0% linked")
    kpi_tile(s, MARGIN, 3.60, 2.55, 1.50, "0.87", "%", "Change failure rate",
             "DIRECTIONAL", AMBER, "3 of 345 deployments")
    kpi_tile(s, 3.34, 3.60, 2.55, 1.50, "2.0", "h", "Mean time to restore (median)",
             "INDICATIVE", ROSE, "n = 3 · all at 2.03 h")

    text(s, MARGIN, 5.32, 5.3, 0.22, "What this source cannot support", size=9.5, color=TEXT,
         bold=True, font=FONT_SEMI)
    limits = ["No deployment event: no environment, release marker or deployment-to-incident link.",
              "No commit-to-build link for one project, so its lead time cannot exist.",
              "No project-month where all four metrics are observed — 4 of 9 rows are partial, 5 minimal."]
    yy = 5.60
    for item in limits:
        text(s, MARGIN, yy, 5.28, 0.34, "·  " + item, size=7.6, color=TEXT_2, line=1.12)
        yy += 0.36

    # dashboard
    shot_x, shot_y, shot_w = 6.30, 1.98, 6.41
    shot_h = shot_w * 1114 / 2360
    pic = s.shapes.add_picture(str(ASSETS / "dashboard-framed.png"), Inches(shot_x), Inches(shot_y),
                               width=Inches(shot_w), height=Inches(shot_h))
    shadow(pic, 26, 10, 5400000, "000000", 55)
    text(s, 6.30, shot_y + shot_h + 0.08, 6.41, 0.2,
         "Live Power BI report — 2 pages, 19 visuals, refresh-verified in Desktop",
         size=7.2, color=MUTED, align="c")

    card(s, 6.30, 5.30, 6.41, 1.42, tone=CARD, alpha=65, shade=None, border_alpha=65)
    text(s, 6.50, 5.42, 6.0, 0.2, "WHAT BECOMES ANSWERABLE AFTER THE MEASUREMENT REBUILD", size=6.8,
         color=CYAN, bold=True, font=FONT_SEMI, letter=1.2)
    after = [("How often do we deploy?", "A discrete event, with environment"),
             ("What is our failure rate?", "Linked to the failing deployment"),
             ("Is a team improving?", "Paired metrics, published sample sizes")]
    for i, (q, a) in enumerate(after):
        x = 6.50 + i * 2.06
        text(s, x, 5.68, 1.95, 0.2, q, size=7.4, color=TEXT, bold=True, font=FONT_SEMI)
        text(s, x, 5.90, 1.95, 0.45, a, size=6.8, color=TEXT_2, line=1.1)

    text(s, 6.50, 6.34, 6.0, 0.2,
         "All four depend on one addition: a deployment event emitted by the golden path.",
         size=6.8, color=MUTED, italic=True)

    footer(s, 7)
    notes(s, "These are the four figures the programme can defend today, each carrying its own confidence "
             "label. The table on the left is the honest limit: no deployment event exists in the source, so "
             "frequency and failure rate are proxies. The screenshot is the working report, not a mock-up. "
             "The rebuild converts proxies into events.")


# ---------------------------------------------------------------------------
# slide 8 - operating model
# ---------------------------------------------------------------------------

def slide_operating_model(prs):
    s = add_slide(prs)
    header(s, 8, "Part C · Strategy and leadership",
           "Operating model", "named owners, written decision rights",
           "Central platform engineering sets the path; application teams own their services; security enforces "
           "in the pull request; data owns the number. Capability gaps close by month 3.", MINT)

    owners = [("Platform product owner", BLUE, "Catalogue, golden path, runner strategy — and the authority to say no"),
              ("Security engineering", ROSE, "Policy-as-code and the exception register — one dedicated engineer, ring-fenced"),
              ("Data and measurement", VIOLET, "Metric definitions, reconciliation gate and the single Gold source"),
              ("Application teams", MINT, "Service quality, tests and on-call — on the golden path, not around it")]
    text(s, MARGIN, 1.98, 3.0, 0.2, "WHO OWNS WHAT", size=7.4, color=MUTED, bold=True,
         font=FONT_SEMI, letter=1.4)
    yy = 2.24
    for role, color, desc in owners:
        card(s, MARGIN, yy, 6.38, 0.62, tone=CARD, alpha=92, shade=None)
        shape(s, MSO_SHAPE.RECTANGLE, MARGIN, yy, 0.05, 0.62, color, alpha=95)
        text(s, MARGIN + 0.20, yy, 2.05, 0.62, role, size=8.8, color=TEXT, bold=True,
             font=FONT_SEMI, anchor="m")
        text(s, MARGIN + 2.28, yy, 3.95, 0.62, desc, size=7.4, color=TEXT_2, line=1.1,
             anchor="m")
        yy += 0.70

    text(s, 7.28, 1.98, 5.0, 0.2, "DECISION RIGHTS — SO THE PROGRAMME STOPS ESCALATING", size=7.4,
         color=MUTED, bold=True, font=FONT_SEMI, letter=1.2)
    rights = [("Approve a new golden-path pattern", "Head of DevOps"), ("Grant a security exception", "CISO"),
              ("Change a published metric definition", "Head of DevOps + Data"),
              ("Re-platform or retire a legacy product", "CTO")]
    yy = 2.24
    for what, who in rights:
        hairline(s, 7.28, yy, 5.45, alpha=30)
        text(s, 7.28, yy + 0.10, 3.55, 0.2, what, size=8.2, color=TEXT_2)
        text(s, 9.05, yy + 0.10, 3.65, 0.2, who, size=8.2, color=CYAN, bold=True, font=FONT_SEMI,
             align="r")
        yy += 0.44
    hairline(s, 7.28, yy, 5.45, alpha=30)
    text(s, 7.28, yy + 0.12, 5.4, 0.44,
         "One named role decides. A committee is not an answer when the CTO asks who approved a control.",
         size=7.4, color=MUTED, italic=True, line=1.15)

    # service tiers
    text(s, MARGIN, 5.24, 3.6, 0.2, "SERVICE TIERS — COMMITMENTS WE CAN KEEP", size=7.4, color=MUTED,
         bold=True, font=FONT_SEMI, letter=1.2)
    tiers = [("T0 · Revenue critical", "Golden path, change freezes", "24×7 on-call · 15-min platform response", MINT),
             ("T1 · Business critical", "Golden path, 99.5% availability", "Extended hours · 1-hour response", CYAN),
             ("T2 · Standard", "Golden path, business hours", "Planned work", BLUE_LT),
             ("Exception", "Documented workaround, owner named", "Monthly review, dated expiry", AMBER)]
    for i, (name, commitment, support, color) in enumerate(tiers):
        x = MARGIN + i * 3.05
        card(s, x, 5.50, 2.92, 1.14, tone=CARD, alpha=90, shade=None)
        shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x + 0.16, 5.62, 0.28, 0.045, color, radius=0.5)
        text(s, x + 0.16, 5.72, 2.6, 0.2, name, size=8.6, color=TEXT, bold=True, font=FONT_SEMI)
        text(s, x + 0.16, 5.96, 2.6, 0.2, commitment, size=7.2, color=TEXT_2)
        text(s, x + 0.16, 6.18, 2.6, 0.34, support, size=7.0, color=color, line=1.1)

    footer(s, 8)
    notes(s, "The internal developer platform is run as a product with a named owner who can say no to "
             "off-catalogue work. Decision rights are explicit so escalation stops. Service tiers convert "
             "platform commitments into something product teams can plan against; the exception tier is the "
             "only one with an expiry date.")


# ---------------------------------------------------------------------------
# slide 9 - roadmap
# ---------------------------------------------------------------------------

def slide_roadmap(prs):
    s = add_slide(prs)
    header(s, 9, "Part C · Strategy and leadership",
           "Roadmap", "five waves, twelve months, one hard gate",
           "Waves are sequenced by dependency and by evidence. Every wave has an exit criterion that can fail.",
           MINT)

    waves = [
        ("WAVE 0", "Days 1–30", "Baseline the estate",
         "Duplicate-pipeline analysis, runner inventory, control gaps, metric reconciliation",
         "Duplicate share quantified; every legacy runner owned and dated", BLUE_LT),
        ("WAVE 1", "Months 2–3", "Golden path v1",
         "Two highest-volume service patterns; controls live on every new pipeline",
         "10 services on the path; 100% of new pipelines pass every control", CYAN),
        ("WAVE 2", "Months 4–6", "First consolidation",
         "Migrate high-volume, low-exception repos; retire the first ADO organization",
         "First org decommissioned; no DORA regression", MINT),
        ("WAVE 3", "Months 7–9", "Scale and settle",
         "Remaining standard repos; legacy products re-platformed or given a dated decision",
         "80% of repositories migrated; no open-ended exceptions", AMBER),
        ("WAVE 4", "Months 10–12", "One organization",
         "Final consolidation, last ADO decommission, handover to run",
         "One org in use; cost reduction realised; run model signed off", VIOLET),
    ]
    text(s, MARGIN, 1.90, 12.0, 0.2, "DELIVERY WAVES", size=7.2, color=MUTED, bold=True,
         font=FONT_SEMI, letter=1.4)
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 1.82, 2.62, 9.68, 0.035,
          grad=[(0, BLUE, 90), (0.5, CYAN, 90), (1, VIOLET, 90)], angle=0, radius=0.5)

    for i, (wave, window, title, scope, exit_, color) in enumerate(waves):
        x = 0.68 + i * 2.42
        c = x + 1.14
        shape(s, MSO_SHAPE.OVAL, c - 0.09, 2.55, 0.18, 0.18, INK, line=color, lw=1.5)
        text(s, x, 2.06, 2.28, 0.2, wave, size=8.2, color=color, bold=True, font=FONT_SEMI, letter=1.3)
        text(s, x, 2.26, 2.28, 0.2, window, size=7.0, color=MUTED)
        card(s, x, 2.86, 2.28, 1.94, tone=CARD, alpha=92, shade=None)
        text(s, x + 0.16, 2.98, 2.0, 0.24, title, size=9.4, color=TEXT, bold=True, font=FONT_SEMI)
        text(s, x + 0.16, 3.26, 1.98, 0.68, scope, size=7.2, color=TEXT_2, line=1.14)
        hairline(s, x + 0.16, 4.02, 1.96, alpha=30)
        text(s, x + 0.16, 4.10, 1.96, 0.16, "EXIT CRITERION", size=5.8, color=color, bold=True,
             font=FONT_SEMI, letter=1.0)
        text(s, x + 0.16, 4.28, 1.96, 0.46, exit_, size=7.0, color=TEXT_2, line=1.1)

    # day-90 gate
    gate = card(s, MARGIN, 4.98, SLIDE_W - 2 * MARGIN, 1.42, tone="12203A", alpha=95,
                border=AMBER, border_alpha=45, shade=(28, 10, 5400000, "000000", 50))
    text(s, MARGIN + 0.24, 5.12, 4.6, 0.22, "DAY-90 GATE — CONTINUE ONLY IF ALL FIVE HOLD", size=7.4,
         color=AMBER, bold=True, font=FONT_SEMI, letter=1.3)
    text(s, MARGIN + 0.24, 5.34, 2.7, 0.72,
         "A year-long plan with no intermediate gate is a budget request, not a plan.",
         size=8.0, color=TEXT_2, line=1.2)

    gates = [("≥ 10 services", "on the golden path"), ("100%", "of new pipelines controlled"),
             ("Every figure", "recomputes green"), ("≥ 1 org", "decommission-ready"),
             ("No regression", "in median lead time")]
    for i, (big, small) in enumerate(gates):
        x = 3.42 + i * 1.86
        shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, 5.42, 1.76, 0.70, "0B1626", alpha=100,
              line=STROKE, lw=0.75, radius=0.1)
        text(s, x + 0.12, 5.50, 1.52, 0.22, big, size=8.4, color=TEXT, bold=True, font=FONT_SEMI,
             align="c")
        text(s, x + 0.12, 5.74, 1.52, 0.28, small, size=6.4, color=MUTED, align="c", line=1.0)

    text(s, MARGIN, 6.54, 7.6, 0.24,
         "Staffing: 4 platform · 1 security · 1 data · 2 migration, plus product owners per wave",
         size=7.2, color=TEXT_2)
    text(s, 7.9, 6.54, 4.83, 0.24,
         "M3 achieved today — the baseline is already published from one Gold source",
         size=7.2, color=MINT, align="r")

    footer(s, 9)
    notes(s, "Five waves, each with a falsifiable exit criterion. The Day-90 gate is the important governance "
             "mechanism: if adoption, control coverage, reproducibility, migration progress or developer friction "
             "miss the threshold, we change the approach while it is still cheap. M3 is already achieved.")


# ---------------------------------------------------------------------------
# slide 10 - risks and stop-list
# ---------------------------------------------------------------------------

def slide_risks(prs):
    s = add_slide(prs)
    header(s, 10, "Part C · Strategy and leadership",
           "Risks", "capacity is the constraint; here is the stop-list",
           "One small team is being asked to migrate an estate, build a platform and fix the measurement. "
           "Something has to come off that team's plate.", AMBER)

    cols = [(0.0, 3.12, "RISK"), (3.12, 2.62, "EARLY WARNING"), (5.74, 1.72, "HELD BY")]
    for dx, w, label in cols:
        text(s, MARGIN + dx + 0.16, 1.96, w, 0.2, label, size=6.6, color=MUTED, bold=True,
             font=FONT_SEMI, letter=1.3)

    risks = [
        ("R5", "One team migrates, builds and measures", "Golden-path backlog age > 2 weeks", "CTO", True),
        ("R1", "Golden path not adopted; teams fork pipelines", "Fork count rising week on week", "Platform owner", False),
        ("R2", "Enforcement pushes delivery off-platform", "Exception volume rising", "CISO", False),
        ("R3", "Legacy products block decommissioning", "Exceptions past their expiry", "CTO", False),
        ("R4", "Metrics stay disputed; progress unprovable", "Dashboard divergence reports", "Head of DevOps", False),
    ]
    yy = 2.18
    for rid, risk, warning, owner, hot in risks:
        card(s, MARGIN, yy, 7.46, 0.60, tone="17223A" if hot else CARD, alpha=95,
             border=AMBER if hot else STROKE, border_alpha=55 if hot else 100, shade=None)
        if hot:
            shape(s, MSO_SHAPE.RECTANGLE, MARGIN, yy, 0.055, 0.60, AMBER, alpha=95)
        text(s, MARGIN + 0.16, yy + 0.19, 0.38, 0.2, rid, size=7.8, color=AMBER if hot else MUTED,
             bold=True, font=FONT_SEMI)
        text(s, MARGIN + 0.58, yy + 0.10, 2.5, 0.44, risk, size=8.2, color=TEXT if hot else TEXT_2,
             line=1.08)
        text(s, MARGIN + 3.28, yy + 0.10, 2.4, 0.44, warning, size=7.6, color=TEXT_2, line=1.08)
        text(s, MARGIN + 5.90, yy + 0.19, 1.42, 0.2, owner, size=7.4, color=CYAN, bold=True,
             font=FONT_SEMI, align="r")
        yy += 0.68

    text(s, MARGIN, 5.70, 7.4, 0.2, "What would make us change the recommendation", size=8.6,
         color=TEXT, bold=True, font=FONT_SEMI)
    changes = ["Golden path rejected by more than a third of pilot teams → pause and re-scope the product",
               "Security exception volume grows week on week → fix the control, not the teams",
               "Migration stalls for two waves → consolidate to three organizations first"]
    yy = 5.96
    for c in changes:
        text(s, MARGIN, yy, 7.42, 0.2, "·  " + c, size=7.2, color=TEXT_2, line=1.05)
        yy += 0.26

    # stop list
    card(s, 8.34, 1.98, 4.37, 2.94, tone="1B1730", alpha=95, border=VIOLET, border_alpha=40)
    text(s, 8.54, 2.12, 3.9, 0.2, "WHAT WE STOP, TO MAKE ROOM", size=7.4, color=VIOLET, bold=True,
         font=FONT_SEMI, letter=1.4)
    stops = [("Per-repository pipeline porting outside Wave 1 pilots",
              "The single largest source of avoidable work"),
             ("Team DORA dashboards that cannot reconcile to Gold",
              "Every unreconciled view restarts the dispute"),
             ("Quarterly audits as the primary control",
              "Enforcement moves into the pull request")]
    yy = 2.44
    for title, why in stops:
        badge(s, 8.54, yy + 0.03, 0.22, "✕", color=ROSE, size=6.5, text_color=INK)
        text(s, 8.86, yy, 3.6, 0.24, title, size=8.2, color=TEXT, bold=True, font=FONT_SEMI,
             line=1.05)
        text(s, 8.86, yy + 0.40, 3.6, 0.2, why, size=7.0, color=MUTED)
        yy += 0.78

    card(s, 8.34, 5.06, 4.37, 1.62, tone=CARD, alpha=90, shade=None)
    text(s, 8.54, 5.18, 3.9, 0.2, "RISK APPETITE, STATED PLAINLY", size=7.2, color=AMBER, bold=True,
         font=FONT_SEMI, letter=1.3)
    appetite = [("Security", "Very low — no new pipeline without enforced controls"),
                ("Delivery pace", "Moderate — waves slip before controls slip"),
                ("Exceptions", "Low and time-bound — no expiry is a failed decision"),
                ("Metric precision", "High — publish blank rather than wrong")]
    yy = 5.42
    for name, statement in appetite:
        text(s, 8.54, yy, 0.95, 0.2, name, size=7.0, color=TEXT_2, bold=True, font=FONT_SEMI)
        text(s, 9.50, yy, 3.05, 0.2, statement, size=6.9, color=TEXT_2)
        yy += 0.28

    footer(s, 10)
    notes(s, "Capacity is the binding constraint: the same team is migrating, building and measuring. The "
             "stop-list is the specific proposal to create room, and it is a decision the CTO can take this "
             "month. Each risk has an early warning that can be seen before the impact and a named holder.")


# ---------------------------------------------------------------------------
# slide 11 - recommendation
# ---------------------------------------------------------------------------

def slide_recommendation(prs):
    s = add_slide(prs)
    header(s, 11, "Part D · Executive recommendation",
           "The choice", "re-scope, do not accelerate or pause",
           "Acceleration scales the defect; a pause stops the only workstream that is working. Re-scoping moves "
           "effort from per-repository conversion to a reusable platform.", MINT)

    options = [
        ("ACCELERATE", "REJECT", ROSE,
         "Per-pipeline conversion continues. Duplication is carried forward at 1,500× cost and enforcement is "
         "deferred again.", "Duplication at 1,500× cost", False),
        ("PAUSE", "REJECT", ROSE,
         "Proven migration stops. Teams keep shipping on inconsistent pipelines; restart cost and credibility "
         "loss are both high.", "Restart and lost credibility", False),
        ("RE-SCOPE", "RECOMMENDED", MINT,
         "Migration continues; conversion becomes a golden path plus a factory; controls apply to new work now; "
         "funding is gated at Day 90.", "Teams give up bespoke pipelines", True),
    ]
    for i, (name, verdict, color, body, cost, hot) in enumerate(options):
        x = MARGIN + i * 4.07
        if hot:
            shp = shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, 1.94, 3.95, 1.90,
                        grad=[(0, "10243A", 100), (1, "122B33", 100)], angle=90,
                        line=MINT, lw=1.25, radius=0.05, shade=(30, 12, 5400000, "000000", 55))
            from presentation_kit import glow
            glow(shp, 14, MINT, 22)
        else:
            card(s, x, 1.98, 3.95, 1.82, tone=CARD, alpha=88, shade=None)
        text(s, x + 0.24, 2.12, 2.7, 0.2, name, size=14, color=TEXT, bold=True, font=FONT_SEMI,
             letter=0.6)
        pill(s, x + 2.82, 2.14, 0.94, 0.24, verdict, color, size=6.2, fill_alpha=14)
        text(s, x + 0.24, 2.48, 3.48, 0.72, body, size=8.0, color=TEXT_2, line=1.22)
        hairline(s, x + 0.24, 3.30, 3.47, alpha=28)
        text(s, x + 0.24, 3.42, 3.47, 0.2,
             [("COST  ", {"size": 6.6, "color": MUTED, "bold": True, "letter": 1.2}),
              (cost, {"size": 7.5, "color": color})], line=1.05)

    columns = [
        ("INVEST NOW", MINT, ["Golden path and migration factory",
                              "Policy-as-code and the exception register",
                              "Measurement spine and reconciliation gate",
                              "Legacy runner isolation, time-boxed"]),
        ("STOP", ROSE, ["Per-repository porting outside pilots",
                        "Team dashboards that cannot reconcile to Gold",
                        "Quarterly audits as the primary control"]),
        ("DEFER", MUTED, ["Azure Boards retirement",
                          "Full ADO decommission before wave exit",
                          "Autonomous production change",
                          "Backfilling historical deployment events"]),
    ]
    for i, (title, color, items) in enumerate(columns):
        x = MARGIN + i * 4.07
        card(s, x, 4.14, 3.95, 2.04, tone=CARD, alpha=88, shade=None)
        text(s, x + 0.22, 4.28, 3.5, 0.2, title, size=7.4, color=color, bold=True, font=FONT_SEMI,
             letter=1.5)
        yy = 4.54
        for item in items:
            text(s, x + 0.22, yy, 0.16, 0.2, "—", size=7.4, color=color)
            text(s, x + 0.44, yy, 3.35, 0.34, item, size=7.8, color=TEXT_2, line=1.08)
            yy += 0.38

    card(s, MARGIN, 6.32, SLIDE_W - 2 * MARGIN, 0.56, tone="0F1B2E", alpha=95, shade=None,
         border_alpha=75)
    text(s, MARGIN + 0.22, 6.40, 12.0, 0.4,
         [("Investment shape  ", {"color": MUTED, "size": 6.8, "bold": True, "letter": 1.2}),
          ("4 platform FTE reallocated from per-repository conversion  ·  1 security FTE ring-fenced  ·  "
           "1 data FTE part-funded by retiring dashboards  ·  legacy isolation charged to the owning product  ·  "
           "tooling neutral to down", {"color": TEXT_2, "size": 7.2})],
         line=1.15, anchor="m")

    footer(s, 11)
    notes(s, "The recommendation is a re-scope with a cost: teams with unusual needs request a pattern rather "
             "than fork a pipeline, and bespoke dashboards are retired. Investment is largely reallocated rather "
             "than new headcount — one security and one data engineer are the exceptions, both justified by "
             "retired work.")


# ---------------------------------------------------------------------------
# slide 12 - decisions and next 90 days
# ---------------------------------------------------------------------------

def slide_decisions(prs):
    s = add_slide(prs)
    header(s, 12, "Part D · Executive recommendation",
           "Decisions", "six requested, four needed this month",
           "Approve D1, D2, D3 and D6 now so Wave 1 starts on a golden path, with enforceable controls and one "
           "agreed number.", MINT)

    text(s, MARGIN, 1.96, 7.4, 0.2, "DECISION", size=6.8, color=MUTED, bold=True, font=FONT_SEMI,
         letter=1.3)
    text(s, 5.36, 1.96, 1.55, 0.2, "OWNER", size=6.8, color=MUTED, bold=True, font=FONT_SEMI,
         letter=1.3)
    text(s, 6.98, 1.96, 1.0, 0.2, "BY", size=6.8, color=MUTED, bold=True, font=FONT_SEMI,
         letter=1.3)

    decisions = [
        ("D1", "Adopt golden-path-first conversion; freeze porting outside pilots", "CTO", "Month 1", True),
        ("D2", "Mandate one DORA definition across all organizations", "CTO + Data", "Month 1", True),
        ("D3", "Enforce five non-waivable controls on all new pipelines", "CISO + CTO", "Month 1", True),
        ("D4", "Fund the legacy runner exception path, capped at 12 months", "CTO + Product", "Month 2", False),
        ("D5", "Decide re-platform or retire for unmodernisable products", "CTO", "Month 6", False),
        ("D6", "Approve the stop-list that funds platform capacity", "CTO", "Month 1", True),
    ]
    yy = 2.20
    for did, what, owner, when, now in decisions:
        card(s, MARGIN, yy, 7.46, 0.52, tone=CARD if not now else "14243A", alpha=95,
             border=MINT if now else STROKE, border_alpha=35 if now else 100, shade=None)
        text(s, MARGIN + 0.16, yy + 0.15, 0.42, 0.2, did, size=8.0, color=MINT if now else MUTED,
             bold=True, font=FONT_SEMI)
        text(s, MARGIN + 0.62, yy + 0.09, 4.05, 0.36, what, size=8.0, color=TEXT_2, line=1.05)
        text(s, 5.36, yy + 0.15, 1.5, 0.2, owner, size=7.4, color=TEXT_2)
        pill(s, 6.98, yy + 0.14, 0.92, 0.24, when, MINT if now else MUTED, size=6.8, fill_alpha=12)
        yy += 0.58

    # what you will see
    card(s, 8.34, 1.98, 4.37, 3.40, tone=CARD, alpha=88, shade=None)
    text(s, 8.54, 2.12, 4.0, 0.2, "WHAT YOU WILL SEE", size=7.2, color=CYAN, bold=True,
         font=FONT_SEMI, letter=1.4)
    milestones = [("Week 1", "Golden-path policy published; conversion frozen outside pilots"),
                  ("Weeks 2–4", "Controls in enforce mode; exception register live with owners and dates"),
                  ("Month 2", "DORA baseline published from one Gold source, with its stated limits"),
                  ("Month 3", "Day-90 gate review: adoption, coverage, reproducibility, migration, friction")]
    shape(s, MSO_SHAPE.RECTANGLE, 8.62, 2.50, 0.014, 2.02, CYAN, alpha=35)
    yy = 2.44
    for when, what in milestones:
        shape(s, MSO_SHAPE.OVAL, 8.57, yy + 0.05, 0.10, 0.10, CYAN, alpha=90)
        text(s, 8.82, yy - 0.02, 1.0, 0.2, when, size=7.2, color=CYAN, bold=True, font=FONT_SEMI)
        text(s, 9.72, yy - 0.02, 2.8, 0.5, what, size=7.2, color=TEXT_2, line=1.12)
        yy += 0.66

    # closing band
    band = shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, MARGIN, 5.72, SLIDE_W - 2 * MARGIN, 1.06,
                 grad=[(0, "12305A", 100), (0.55, "2A1A5E", 100), (1, "123A44", 100)], angle=0,
                 line="3A5C8F", lw=0.75, radius=0.06, shade=(30, 10, 5400000, "000000", 50))
    text(s, MARGIN + 0.28, 5.84, 1.8, 0.2, "THE ASK", size=7.2, color=MINT, bold=True,
         font=FONT_SEMI, letter=1.8)
    text(s, MARGIN + 0.28, 6.04, 9.4, 0.62,
         "Approve D1, D2, D3 and D6 this month. Everything else in this brief follows from those four decisions.",
         size=12.4, color=TEXT, bold=True, font=FONT_SEMI, line=1.15)
    text(s, 10.6, 6.10, 1.9, 0.5,
         "Head of DevOps\n& Analytics\n25 October 2026", size=7.1, color=TEXT_2, align="r", line=1.25)

    footer(s, 12)
    notes(s, "Close on the four decisions needed this month: golden-path-first, one metric definition, enforced "
             "controls, and the stop-list that funds the team. Each has an owner, a date and a visible artefact. "
             "The Day-90 gate is the point at which the CTO sees evidence before releasing full funding.")


# ---------------------------------------------------------------------------
# slide 13 - questions and discussion
# ---------------------------------------------------------------------------

def slide_qa(prs):
    s = add_slide(prs, "cover")

    logo(s, 0.88, 0.62, 0.40)
    text(s, 1.23, 1.06, 2.2, 0.2, "Make it amazing", size=7.5, color=TEXT_2, letter=0.6)

    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.88, 2.02, 0.34, 0.028, BLUE, radius=0.5)
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 1.24, 2.02, 0.34, 0.028, VIOLET, radius=0.5)
    text(s, 1.74, 1.92, 7.2, 0.25, "CLOSING  ·  DISCUSSION", size=9, color=CYAN,
         bold=True, font=FONT_SEMI, letter=2.2)

    text(s, 0.88, 2.34, 6.4, 1.55, "Questions &\ndiscussion.", size=40, color=TEXT, bold=True,
         font=FONT_SEMI, line=1.04)

    text(s, 0.90, 4.10, 5.9, 0.72,
         "Thank you. The strategy, the working pipeline and every figure in this deck are in the "
         "repository, so the debate can move from opinion to evidence.",
         size=10.5, color=TEXT_2, line=1.35)

    hairline(s, 0.88, 5.12, 5.55, alpha=45)

    text(s, 0.88, 5.28, 4.6, 0.22, "Head of DevOps & Analytics", size=10, color=TEXT, bold=True,
         font=FONT_SEMI)
    text(s, 0.88, 5.54, 5.4, 0.4,
         "Prepared for the CTO and Technology Leadership Team\n25 October 2026  ·  Internal — for decision",
         size=8.2, color=MUTED, line=1.25)

    text(s, 0.88, 6.42, 5.6, 0.2, "DECISIONS ON THE TABLE", size=7.2, color=MUTED, bold=True,
         font=FONT_SEMI, letter=1.4)
    recap = [("D1", "Golden path", BLUE_LT), ("D2", "One definition", CYAN),
             ("D3", "Controls", VIOLET), ("D6", "Stop-list", MINT)]
    for i, (did, label, color) in enumerate(recap):
        x = 0.88 + i * 1.44
        pill(s, x, 6.66, 1.34, 0.28, f"{did}  ·  {label}", color, size=6.8, fill_alpha=11)

    # discussion prompts
    prompts = [
        ("SCOPE", "Which services enter Wave 1 — and who owns the pilot outcome?",
         "The pilot sets the evidence for everything that follows.", BLUE_LT),
        ("CONSTRAINTS", "Which legacy products need a dated runner exception, and until when?",
         "Every exception needs a named owner and an expiry date.", AMBER),
        ("EVIDENCE", "What must you see at Day 90 to release full-estate funding?",
         "The gate is the release valve, not a formality.", MINT),
    ]
    for i, (label, question, hint, color) in enumerate(prompts):
        y = 2.02 + i * 1.46
        card(s, 7.20, y, 5.51, 1.25, tone=CARD, alpha=88, shade=None)
        shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 7.20, y, 0.055, 1.25, color, alpha=95, radius=0.5)
        text(s, 7.44, y + 0.16, 4.9, 0.2, label, size=7.2, color=color, bold=True,
             font=FONT_SEMI, letter=1.6)
        text(s, 7.44, y + 0.40, 4.95, 0.34, question, size=10.5, color=TEXT, line=1.15)
        text(s, 7.44, y + 0.86, 4.95, 0.24, hint, size=7.6, color=MUTED, italic=True)

    text(s, 7.20, 6.48, 5.5, 0.2, "The answer to each is a decision, not a workstream.",
         size=7.6, color=MUTED, italic=True, align="r")

    text(s, 11.3, 7.10, 1.4, 0.2, f"{TOTAL:02d} / {TOTAL:02d}", size=7.5, color=MUTED, align="r")

    notes(s, "Close the session by inviting the three decisions that unlock the next wave: scope of the pilot, "
             "the dated exceptions for legacy products, and the evidence required at Day 90. Keep the "
             "discussion anchored to the four decisions requested on the previous slide.")


# ---------------------------------------------------------------------------

def main():
    build_assets()
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)

    slide_cover(prs)
    slide_summary(prs)
    slide_architecture(prs)
    slide_golden_path(prs)
    slide_ai(prs)
    slide_data_spine(prs)
    slide_metrics(prs)
    slide_operating_model(prs)
    slide_roadmap(prs)
    slide_risks(prs)
    slide_recommendation(prs)
    slide_decisions(prs)
    slide_qa(prs)

    for slide in prs.slides:
        add_transition(slide, "fade", "med")

    prs.core_properties.title = "Cyncly Engineering Transformation — CTO Briefing"
    prs.core_properties.author = "Head of DevOps & Analytics"
    prs.core_properties.subject = "Consolidating five Azure DevOps organizations onto GitHub Enterprise"
    prs.core_properties.keywords = "Cyncly, DevOps, GitHub Enterprise, DORA, Power BI, strategy"
    prs.core_properties.comments = "Internal — for decision."

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    print(f"wrote {OUT} ({OUT.stat().st_size / 1_000_000:.1f} MB, {len(prs.slides)} slides)")


if __name__ == "__main__":
    main()
